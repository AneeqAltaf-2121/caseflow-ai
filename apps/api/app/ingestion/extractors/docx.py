import io
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.errors import ValidationError
from app.ingestion.models import ExtractedDocument, ExtractedPage


class DocxExtractor:
    """Modern .docx (OOXML) only — legacy binary .doc isn't parseable by
    python-docx and isn't in DocumentService.ALLOWED_CONTENT_TYPES for
    that reason (see that module's comment).

    .docx has no reliable page boundaries in the file format itself (page
    breaks depend on rendering), so the whole document extracts as a
    single page — paragraph breaks are preserved for chunking (Phase 10)
    to work with instead.
    """

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument:
        try:
            document = DocxDocument(io.BytesIO(data))
        except (PackageNotFoundError, BadZipFile) as exc:
            raise ValidationError(f"Could not read DOCX {filename!r}: {exc}") from exc

        text = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
        return ExtractedDocument(
            source_filename=filename, pages=[ExtractedPage(page_number=1, text=text)]
        )
